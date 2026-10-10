#!/usr/bin/env python3
"""Prepare a private, atomic ACC classification snapshot from the supplied workbook.

The generated SQL contains barcode/classification data and must never be committed.
Run it with ``psql -v ON_ERROR_STOP=1 -f <generated.sql>`` only after the schema
migration is deployed. The active snapshot changes in the same transaction as COPY.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import os
import re
import sys
import unicodedata
import uuid
from collections import Counter, defaultdict
from pathlib import Path


BARCODE = re.compile(r"(?<!\d)\d{8,14}(?!\d)")
PLACEHOLDERS = {"", "0", "не привязано"}
SQL_COLUMNS = (
    "snapshot_id", "ware_id", "barcode", "group_key", "group_label",
    "subgroup_key", "subgroup_label", "mnn_key", "mnn_label",
)


def clean_label(value: object) -> str:
    if value is None:
        return ""
    text = str(value).replace("_x000D_", " ").replace("\x00", " ")
    return " ".join(unicodedata.normalize("NFKC", text).split())


def meaningful_label(value: object) -> str:
    label = clean_label(value)
    return "" if label.casefold() in PLACEHOLDERS else label


def normalized(value: str) -> str:
    return " ".join(unicodedata.normalize("NFKC", value).lower().split())


def scope_key(kind: str, label: str, parent: str = "") -> str:
    prefix, namespace = {
        "group": ("grp_", "acc_group"),
        "subgroup": ("sub_", "acc_subgroup"),
        "mnn": ("mnn_", "acc_mnn"),
    }[kind]
    parts = [namespace]
    if kind == "subgroup":
        if not parent:
            raise ValueError("A subgroup needs a non-empty parent group")
        parts.append(normalized(parent))
    parts.append(normalized(label))
    digest = hashlib.sha256("\x1f".join(parts).encode("utf-8")).hexdigest()[:32]
    return prefix + digest


def barcodes(value: object) -> list[str]:
    if value is None:
        return []
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return list(dict.fromkeys(BARCODE.findall(str(value))))


def sql_literal(value: str) -> str:
    return "'" + value.replace("'", "''") + "'"


def prepare_rows(rows: object, snapshot_id: str) -> tuple[list[tuple[str | None, ...]], dict]:
    """Consume workbook rows after the header; omit only unusable identifiers.

    Every distinct WARE_ID/barcode remains in the snapshot. The backend deliberately
    rejects a barcode shared by different WARE_IDs instead of guessing a medicine.
    """
    prepared: dict[tuple[str, str], tuple[str | None, ...]] = {}
    seen_barcodes: dict[str, set[str]] = defaultdict(set)
    stats: Counter[str] = Counter()
    for row in rows:
        if not any(cell is not None for cell in row):
            continue
        stats["source_rows"] += 1
        ware_id = clean_label(row[0] if len(row) > 0 else None)
        if not ware_id or len(ware_id) > 64:
            stats["missing_or_invalid_ware_id"] += 1
            continue
        found_barcodes = barcodes(row[3] if len(row) > 3 else None)
        if not found_barcodes:
            stats["no_valid_barcode"] += 1
            continue

        group = meaningful_label(row[19] if len(row) > 19 else None)
        subgroup = meaningful_label(row[20] if len(row) > 20 else None)
        mnn = meaningful_label(row[15] if len(row) > 15 else None)
        if subgroup and not group:
            stats["subgroup_without_parent"] += 1
            subgroup = ""
        if not (group or mnn):
            stats["no_usable_scope"] += 1
        for label in (group, subgroup, mnn):
            if len(label) > 255:
                raise ValueError(f"ACC classification label exceeds 255 characters: {label[:80]!r}")

        group_key = scope_key("group", group) if group else None
        subgroup_key = scope_key("subgroup", subgroup, group) if subgroup else None
        mnn_key = scope_key("mnn", mnn) if mnn else None
        for barcode in found_barcodes:
            key = (ware_id, barcode)
            record = (
                snapshot_id, ware_id, barcode, group_key, group or None,
                subgroup_key, subgroup or None, mnn_key, mnn or None,
            )
            if key in prepared:
                if prepared[key] != record:
                    raise ValueError(f"Conflicting classification for WARE_ID/barcode {key!r}")
                stats["identical_duplicate_rows"] += 1
                continue
            prepared[key] = record
            seen_barcodes[barcode].add(ware_id)

    result = sorted(prepared.values(), key=lambda r: (r[2] or "", r[1] or ""))
    stats["mapped_rows"] = len(result)
    stats["distinct_barcodes"] = len(seen_barcodes)
    stats["ambiguous_barcodes"] = sum(len(ware_ids) > 1 for ware_ids in seen_barcodes.values())
    stats["unambiguous_barcodes"] = stats["distinct_barcodes"] - stats["ambiguous_barcodes"]
    return result, dict(sorted(stats.items()))


def write_private(path: Path, content: str) -> None:
    descriptor = os.open(path, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
    try:
        with os.fdopen(descriptor, "w", encoding="utf-8", newline="") as out:
            out.write(content)
    except BaseException:
        path.unlink(missing_ok=True)
        raise


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("output_dir", type=Path)
    args = parser.parse_args()
    if not args.workbook.is_file():
        parser.error("workbook does not exist")
    if args.output_dir.exists():
        parser.error("output directory must not already exist")

    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise SystemExit("openpyxl is required to read the ACC workbook") from exc

    workbook_bytes = args.workbook.read_bytes()
    source_sha256 = hashlib.sha256(workbook_bytes).hexdigest()
    workbook = load_workbook(args.workbook, read_only=True, data_only=True)
    if "Лист1" not in workbook.sheetnames:
        raise SystemExit("Expected ACC sheet 'Лист1' is missing")
    sheet = workbook["Лист1"]
    sheet_rows = sheet.iter_rows(values_only=True)
    header = next(sheet_rows)
    required = {0: "WARE_ID", 3: "ШТРИХКОД", 19: "ГРУППА", 20: "ПОДГРУППА"}
    for index, expected in required.items():
        if clean_label(header[index]) != expected:
            raise SystemExit(f"Column {index + 1} must be {expected!r}")
    if clean_label(header[15]) != "МНН":
        raise SystemExit("Column 16 must be 'МНН'")

    snapshot_id = str(uuid.uuid5(uuid.NAMESPACE_URL, f"epharm/acc-catalog/{source_sha256}"))
    records, stats = prepare_rows(sheet_rows, snapshot_id)
    workbook.close()
    if not records:
        raise SystemExit("No usable ACC barcode classifications found")

    output = io.StringIO()
    output.write("-- Private ACC classification snapshot. Do not commit or paste into logs.\n")
    output.write("BEGIN;\n")
    output.write(
        "CREATE TEMP TABLE acc_catalog_import (LIKE acc_catalog_barcodes INCLUDING DEFAULTS) "
        "ON COMMIT DROP;\n"
    )
    output.write(
        "COPY acc_catalog_import (" + ", ".join(SQL_COLUMNS) + ") "
        "FROM STDIN WITH (FORMAT csv, DELIMITER E'\\t', NULL '\\N');\n"
    )
    writer = csv.writer(output, delimiter="\t", lineterminator="\n")
    for record in records:
        writer.writerow([value if value is not None else r"\N" for value in record])
    output.write("\\.\n")
    output.write(
        "INSERT INTO acc_catalog_snapshots "
        "(id, sha256, source_name, item_count, barcode_count, imported_at) VALUES ("
        f"{sql_literal(snapshot_id)}, {sql_literal(source_sha256)}, "
        f"{sql_literal(args.workbook.name)}, {stats['source_rows']}, "
        f"{stats['distinct_barcodes']}, now()) "
        "ON CONFLICT (sha256) DO NOTHING;\n"
    )
    output.write(
        "INSERT INTO acc_catalog_barcodes (" + ", ".join(SQL_COLUMNS) + ") "
        "SELECT snapshot.id, imported.ware_id, imported.barcode, imported.group_key, "
        "imported.group_label, imported.subgroup_key, imported.subgroup_label, "
        "imported.mnn_key, imported.mnn_label FROM acc_catalog_import AS imported "
        f"CROSS JOIN (SELECT id FROM acc_catalog_snapshots WHERE sha256 = {sql_literal(source_sha256)}) "
        "AS snapshot WHERE true ON CONFLICT (snapshot_id, ware_id, barcode) DO NOTHING;\n"
    )
    output.write(
        "INSERT INTO acc_catalog_state (singleton, active_snapshot_id) "
        f"VALUES (1, (SELECT id FROM acc_catalog_snapshots WHERE sha256 = {sql_literal(source_sha256)})) "
        "ON CONFLICT (singleton) DO UPDATE SET active_snapshot_id = EXCLUDED.active_snapshot_id;\n"
    )
    output.write("COMMIT;\n")

    args.output_dir.mkdir(parents=True, mode=0o700)
    os.chmod(args.output_dir, 0o700)
    sql_path = args.output_dir / "acc_catalog_snapshot.sql"
    manifest_path = args.output_dir / "manifest.json"
    write_private(sql_path, output.getvalue())
    manifest = {
        "snapshot_id": snapshot_id,
        "source_name": args.workbook.name,
        "source_sha256": source_sha256,
        "sql_sha256": hashlib.sha256(sql_path.read_bytes()).hexdigest(),
        "statistics": stats,
    }
    write_private(manifest_path, json.dumps(manifest, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps({"output_dir": str(args.output_dir), **manifest}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
